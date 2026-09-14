#!/usr/bin/env python3
"""Generate test resume files (PDF + DOCX) for verifying the upload & parse flow."""
import subprocess, sys

PDF_PATH = "/home/z/my-project/scripts/test-resume.pdf"
DOCX_PATH = "/home/z/my-project/scripts/test-resume.docx"

# ---------- PDF via reportlab ----------
try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(PDF_PATH, pagesize=A4)
    w, h = A4
    y = h - 30 * mm
    lines = [
        ("Rahul Malhotra", 16),
        ("Email: rahul.malhotra.dev@gmail.com | Phone: +91 98765 12345", 10),
        ("Location: New Delhi, India", 10),
        ("", 10),
        ("SUMMARY", 12),
        ("Full-stack developer with 5 years of experience building web applications", 10),
        ("using React, TypeScript and Node.js.", 10),
        ("", 10),
        ("SKILLS", 12),
        ("JavaScript, TypeScript, React, Next.js, Node.js, Express, PostgreSQL,", 10),
        ("Docker, AWS, REST API, Git, Jest", 10),
        ("", 10),
        ("EXPERIENCE", 12),
        ("Software Engineer II - TechSolutions Pvt Ltd (2021 - Present)", 10),
        ("- Built customer dashboards with React and TypeScript", 10),
        ("- Designed REST APIs with Node.js and PostgreSQL", 10),
        ("Software Engineer - StartUp Labs (2019 - 2021)", 10),
        ("- Developed e-commerce features used by 50k+ monthly users", 10),
        ("", 10),
        ("EDUCATION", 12),
        ("B.Tech Computer Science, Delhi Technological University, 2019", 10),
    ]
    for text, size in lines:
        c.setFont("Helvetica", size)
        c.drawString(25 * mm, y, text)
        y -= (7 if size == 10 else 9) * mm
    c.save()
    print("PDF written:", PDF_PATH)
except ImportError:
    print("reportlab not available", file=sys.stderr)
    sys.exit(1)

# ---------- DOCX via python-docx ----------
try:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Priya Nair")
    doc.add_paragraph("Email: priya.nair.dev@example.com | Phone: +91 90909 11111")
    doc.add_paragraph("Location: Kochi, India")
    doc.add_paragraph("")
    doc.add_paragraph("SKILLS")
    doc.add_paragraph("React, TypeScript, GraphQL, Node.js, MongoDB, Docker, CI/CD")
    doc.add_paragraph("")
    doc.add_paragraph("EXPERIENCE")
    doc.add_paragraph("Senior Software Engineer - WebWorks (2020 - Present), 5 years")
    doc.add_paragraph("- Led frontend platform team for SaaS product")
    doc.add_paragraph("")
    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("M.C.A., Anna University, 2018")
    doc.save(DOCX_PATH)
    print("DOCX written:", DOCX_PATH)
except ImportError:
    print("python-docx not available (skipping DOCX)", file=sys.stderr)
